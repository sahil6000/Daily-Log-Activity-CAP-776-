import openpyxl
from datetime import datetime, timedelta


# ============================================================
# FILE HANDLING
# ============================================================

file_name = "12604444.xlsx"

try:

    workbook = openpyxl.load_workbook(file_name, data_only=True)

    if "Daily Log" not in workbook.sheetnames:
        print("Error: Daily Log sheet not found.")
        exit()

    sheet = workbook["Daily Log"]

    print("Excel file loaded successfully")
    print("Sheet:", sheet.title)

except FileNotFoundError:

    print("Error: Excel file not found.")
    exit()

except Exception as e:

    print("Error while opening Excel file:", e)
    exit()


# ============================================================
# PROJECT DATE RANGE
# ============================================================

start_date = datetime(2026, 8, 13)
end_date = datetime(2026, 9, 21)


# ============================================================
# READ DAILY RECORDS
# ============================================================

records = []

for row in sheet.iter_rows(min_row=5, values_only=True):

    date = row[0]

    if not isinstance(date, datetime):
        continue

    if start_date <= date <= end_date:
        records.append(row)


# ============================================================
# VALIDATION
# ============================================================

valid_feelings = [
    "Excellent",
    "Good",
    "Neutral",
    "Low",
    "Stressed"
]

valid_satisfaction = [
    "Very Satisfied",
    "Satisfied",
    "Neutral",
    "Unsatisfied",
    "Very Unsatisfied"
]

valid_energy = [
    "High",
    "Medium",
    "Low"
]


def validate_records(records):

    valid_records = []
    invalid_records = []

    for row in records:

        valid = True

        # Check numeric columns
        for value in row[1:10]:

            if not isinstance(value, (int, float)):
                valid = False

        # Check Day's Feeling
        if row[10] not in valid_feelings:
            valid = False

        # Check Satisfaction Level
        if row[11] not in valid_satisfaction:
            valid = False

        # Check Energy Level
        if row[12] not in valid_energy:
            valid = False

        if valid:
            valid_records.append(row)

        else:
            invalid_records.append(row)

    return valid_records, invalid_records


valid_records, invalid_records = validate_records(records)


# ============================================================
# SCORE CONVERSION
# ============================================================

def get_feeling_score(feeling):

    if feeling == "Excellent":
        return 5

    elif feeling == "Good":
        return 4

    elif feeling == "Neutral":
        return 3

    elif feeling == "Low":
        return 2

    elif feeling == "Stressed":
        return 1


def get_satisfaction_score(satisfaction):

    if satisfaction == "Very Satisfied":
        return 5

    elif satisfaction == "Satisfied":
        return 4

    elif satisfaction == "Neutral":
        return 3

    elif satisfaction == "Unsatisfied":
        return 2

    elif satisfaction == "Very Unsatisfied":
        return 1


def get_energy_score(energy):

    if energy == "High":
        return 3

    elif energy == "Medium":
        return 2

    elif energy == "Low":
        return 1


# ============================================================
# DATA CONTINUITY CHECK
# ============================================================

def check_date_continuity(valid_records, start_date, end_date):

    recorded_dates = []

    for row in valid_records:
        recorded_dates.append(row[0].date())

    # Check duplicate dates
    duplicate_dates = []

    for date in recorded_dates:

        if recorded_dates.count(date) > 1:
            if date not in duplicate_dates:
                duplicate_dates.append(date)

    # Create expected dates
    expected_dates = []

    current_date = start_date.date()

    while current_date <= end_date.date():

        expected_dates.append(current_date)

        current_date = current_date + timedelta(days=1)

    # Check missing dates
    missing_dates = []

    for date in expected_dates:

        if date not in recorded_dates:
            missing_dates.append(date)

    return missing_dates, duplicate_dates


missing_dates, duplicate_dates = check_date_continuity(
    valid_records,
    start_date,
    end_date
)


# ============================================================
# RESULTS
# ============================================================

expected_days = (end_date - start_date).days + 1

print("\nData Continuity Check:")
print("Missing dates:", len(missing_dates))
print("Duplicate dates:", len(duplicate_dates))

print("\nProject Period:")
print("Start Date:", start_date.date())
print("End Date:", end_date.date())

print("\nExpected days:", expected_days)
print("Records found:", len(records))
print("Valid records:", len(valid_records))
print("Invalid records:", len(invalid_records))


# ============================================================
# ACTIVITY SUMMARY
# ============================================================

def calculate_activity_averages(valid_records):

    total_sleep = 0
    total_fitness = 0
    total_study = 0
    total_coding = 0
    total_class = 0
    total_other = 0
    total_free = 0

    for row in valid_records:

        total_sleep += row[1]
        total_fitness += row[2]
        total_study += row[3]
        total_coding += row[4]
        total_class += row[5]
        total_other += row[7]
        total_free += row[9]

    days = len(valid_records)

    return (
        total_sleep / days,
        total_fitness / days,
        total_study / days,
        total_coding / days,
        total_class / days,
        total_other / days,
        total_free / days
    )


(
    average_sleep,
    average_fitness,
    average_study,
    average_coding,
    average_class,
    average_other,
    average_free
) = calculate_activity_averages(valid_records)


print("\n1. Activity Data Summary:")

print(
    "Average Sleep/day:",
    round(average_sleep, 2),
    "min"
)

print(
    "Average Fitness/day:",
    round(average_fitness, 2),
    "min"
)

print(
    "Average Study/day:",
    round(average_study, 2),
    "min"
)

print(
    "Average Coding/day:",
    round(average_coding, 2),
    "min"
)

print(
    "Average Class/day:",
    round(average_class, 2),
    "min"
)

print(
    "Average Other Activities/day:",
    round(average_other, 2),
    "min"
)

print(
    "Average Free/Unaccounted Time/day:",
    round(average_free, 2),
    "min"
)


# ============================================================
# TPI - TECH PRODUCTIVITY INDEX
# ============================================================

def calculate_tpi(valid_records):

    total_coding = 0

    for row in valid_records:
        total_coding += row[4]

    return total_coding / len(valid_records)


tpi = calculate_tpi(valid_records)


# ============================================================
# AAI - ACADEMIC ACTIVITY INDEX
# ============================================================

def calculate_aai(valid_records):

    total_academic = 0

    for row in valid_records:
        total_academic += row[3] + row[5]

    return total_academic / len(valid_records)


aai = calculate_aai(valid_records)


# ============================================================
# PhAI - PHYSICAL ACTIVITY INDEX
# ============================================================

def calculate_phai(valid_records):

    total_fitness = 0

    for row in valid_records:
        total_fitness += row[2]

    return total_fitness / len(valid_records)


phai = calculate_phai(valid_records)


# ============================================================
# SRI - SLEEP & RECOVERY INDEX
# ============================================================

def calculate_sri(valid_records):

    total_sleep = 0

    for row in valid_records:
        total_sleep += row[1]

    return total_sleep / len(valid_records)


sri = calculate_sri(valid_records)


# ============================================================
# ABI - ACTIVITY BALANCE INDEX
# ============================================================

def calculate_abi(valid_records):

    total_free_time = 0

    for row in valid_records:
        total_free_time += row[9]

    return total_free_time / len(valid_records)


abi = calculate_abi(valid_records)


# ============================================================
# TUI - TIME UTILIZATION INDEX
# ============================================================

def calculate_tui(valid_records):

    total_tracked_time = 0

    for row in valid_records:
        total_tracked_time += row[8]

    return total_tracked_time / len(valid_records)


tui = calculate_tui(valid_records)


# ============================================================
# EI - EXPERIENCE INDEX
# ============================================================

def calculate_ei(valid_records):

    total_ei = 0

    for row in valid_records:

        feeling = row[10]
        satisfaction = row[11]
        energy = row[12]

        feeling_score = get_feeling_score(feeling)

        satisfaction_score = get_satisfaction_score(
            satisfaction
        )

        energy_score = get_energy_score(energy)

        # Daily EI
        daily_ei = (
            feeling_score
            + satisfaction_score
            + energy_score
        ) / 3

        total_ei += daily_ei

    return total_ei / len(valid_records)


ei = calculate_ei(valid_records)


# ============================================================
# DCI - DATA CONTINUITY INDEX
# ============================================================

dci = (len(valid_records) / expected_days) * 100


# ============================================================
# PAI - PERSONAL ACTIVITY INDEX
# ============================================================

def calculate_pai(
    tpi,
    aai,
    phai,
    sri,
    tui,
    ei,
    dci
):

    pai = (
        (0.15 * tpi)
        + (0.20 * aai)
        + (0.15 * phai)
        + (0.20 * sri)
        + (0.15 * tui)
        + (0.10 * ei)
        + (0.05 * dci)
    )

    return pai


pai = calculate_pai(
    tpi,
    aai,
    phai,
    sri,
    tui,
    ei,
    dci
)


# ============================================================
# CORRELATION FUNCTION
# ============================================================

def calculate_correlation(x_values, y_values):

    n = len(x_values)

    mean_x = sum(x_values) / n
    mean_y = sum(y_values) / n

    numerator = 0
    sum_x = 0
    sum_y = 0

    for i in range(n):

        x_difference = x_values[i] - mean_x
        y_difference = y_values[i] - mean_y

        numerator += (
            x_difference * y_difference
        )

        sum_x += x_difference ** 2
        sum_y += y_difference ** 2

    denominator = (
        sum_x * sum_y
    ) ** 0.5

    if denominator == 0:
        return 0

    return numerator / denominator


# ============================================================
# RELATIONSHIP 1 - SLEEP AND ENERGY
# ============================================================

sleep_values = []
energy_values = []

for row in valid_records:

    sleep = row[1]
    energy = row[12]

    energy_score = get_energy_score(energy)

    sleep_values.append(sleep)
    energy_values.append(energy_score)


sleep_energy_correlation = calculate_correlation(
    sleep_values,
    energy_values
)


# ============================================================
# RELATIONSHIP 2 - STUDY AND SATISFACTION
# ============================================================

study_values = []
satisfaction_values = []

for row in valid_records:

    study = row[3]
    satisfaction = row[11]

    satisfaction_score = get_satisfaction_score(
        satisfaction
    )

    study_values.append(study)
    satisfaction_values.append(
        satisfaction_score
    )


study_satisfaction_correlation = calculate_correlation(
    study_values,
    satisfaction_values
)


# ============================================================
# RELATIONSHIP 3 - CODING AND ENERGY
# ============================================================

coding_values = []
energy_values = []

for row in valid_records:

    coding = row[4]
    energy = row[12]

    energy_score = get_energy_score(energy)

    coding_values.append(coding)
    energy_values.append(energy_score)


coding_energy_correlation = calculate_correlation(
    coding_values,
    energy_values
)


# ============================================================
# FINAL OUTPUT
# ============================================================

print("\n2. Index Values:")

print(
    "TPI:",
    round(tpi, 2),
    "min/day"
)

print(
    "AAI:",
    round(aai, 2),
    "min/day"
)

print(
    "PhAI:",
    round(phai, 2),
    "min/day"
)

print(
    "SRI:",
    round(sri, 2),
    "min/day"
)

print(
    "ABI:",
    round(abi, 2),
    "min/day"
)

print(
    "TUI:",
    round(tui, 2),
    "min/day"
)

print(
    "EI:",
    round(ei, 2)
)

print(
    "DCI:",
    round(dci, 2),
    "%"
)

print(
    "PAI:",
    round(pai, 2)
)

# ============================================================
# KEY FINDINGS
# ============================================================

print("\n3. Key Findings/Correlation from My Data:")

print(
    "Sleep-Energy Correlation:",
    round(sleep_energy_correlation, 2)
)

print(
    "Study-Satisfaction Correlation:",
    round(study_satisfaction_correlation, 2)
)

print(
    "Coding-Energy Correlation:",
    round(coding_energy_correlation, 2)
)

# ============================================================
# ADDITIONAL RELATIONSHIPS
# ============================================================

# ---------------- FEELING - SATISFACTION ----------------

feeling_values = []
satisfaction_values = []

for row in valid_records:

    feeling = row[10]
    satisfaction = row[11]

    feeling_score = get_feeling_score(feeling)
    satisfaction_score = get_satisfaction_score(satisfaction)

    feeling_values.append(feeling_score)
    satisfaction_values.append(satisfaction_score)


feeling_satisfaction = calculate_correlation(
    feeling_values,
    satisfaction_values
)


# ---------------- FEELING - ENERGY ----------------

feeling_values = []
energy_values = []

for row in valid_records:

    feeling = row[10]
    energy = row[12]

    feeling_score = get_feeling_score(feeling)
    energy_score = get_energy_score(energy)

    feeling_values.append(feeling_score)
    energy_values.append(energy_score)


feeling_energy = calculate_correlation(
    feeling_values,
    energy_values
)


# ============================================================
# ADDITIONAL RELATIONSHIPS OUTPUT
# ============================================================

print("\nAdditional Relationships:")

print(
    "Feeling-Satisfaction Correlation:",
    round(feeling_satisfaction, 2)
)

print(
    "Feeling-Energy Correlation:",
    round(feeling_energy, 2)
)